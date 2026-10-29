import streamlit as st
import streamlit.components.v1 as components
import sqlite3
import pandas as pd
from datetime import datetime
import time
import hashlib
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
import os
import glob
import json
import base64
import re
import urllib.request
import pypdf
import io

from setup_inicial import instalar_sistema
from exporter_bys import generar_formulario_oficial_bys
from pdf_signature import exportar_excel_a_pdf, verificar_firma_digital_adobe
from mailer import enviar_correo_con_pdf_adjunto, enviar_correo_pedimento_firmado

EXCEL_MAESTRO = "Control_Maestro_238_Articulos_DINADECO.xlsx"
MAX_INACTIVIDAD = 300  # 5 minutos de inactividad
CLAVE_TEMPORAL_DEFAULT = "Dinadeco2026*"

LISTA_USUARIOS_OFICIALES = [
    {"nombre": "Emmanuel Alvarez", "email": "ealvarez@dinadeco.go.cr", "username": "ealvarez", "puesto": "Misceláneo/a", "espacio": "Piso 3 (30%)", "rol": "BASICO"},
    {"nombre": "Ana Orozco Sánchez", "email": "aorozco@dinadeco.go.cr", "username": "aorozco", "puesto": "Misceláneo/a", "espacio": "Piso 1 (30%)", "rol": "BASICO"},
    {"nombre": "Ana Yancy González", "email": "agonzalez@dinadeco.go.cr", "username": "agonzalez", "puesto": "Recepcionista", "espacio": "Oficina General", "rol": "ADMIN"},
    {"nombre": "Pamela Quirós", "email": "pquiros@dinadeco.go.cr", "username": "pquiros", "puesto": "Administrativa", "espacio": "Oficina General", "rol": "ADMIN"},
    {"nombre": "Glenda Loría", "email": "gloria@dinadeco.go.cr", "username": "gloria", "puesto": "Administrativa", "espacio": "Oficina General", "rol": "ADMIN"},
    {"nombre": "Ronald Montero", "email": "rmontero@dinadeco.go.cr", "username": "rmontero", "puesto": "Chofer", "espacio": "Flota Vehicular", "rol": "BASICO"},
    {"nombre": "Randall Solano", "email": "rsolano@dinadeco.go.cr", "username": "rsolano", "puesto": "Chofer", "espacio": "Flota Vehicular", "rol": "BASICO"},
    {"nombre": "Nelson Levey", "email": "nlevey@dinadeco.go.cr", "username": "nlevey", "puesto": "Chofer", "espacio": "Flota Vehicular", "rol": "BASICO"},
    {"nombre": "Hans Arrieta Morales", "email": "harrieta@dinadeco.go.cr", "username": "harrieta", "puesto": "Oficial de Seguridad", "espacio": "Exteriores", "rol": "BASICO"},
    {"nombre": "Henry Guevara", "email": "hguevara@dinadeco.go.cr", "username": "hguevara", "puesto": "Oficial de Seguridad", "espacio": "Exteriores", "rol": "BASICO"},
    {"nombre": "Jorge Barquero", "email": "jbarquero@dinadeco.go.cr", "username": "jbarquero", "puesto": "Oficial de Seguridad", "espacio": "Exteriores", "rol": "BASICO"},
    {"nombre": "Miguel Granados", "email": "mgranados@dinadeco.go.cr", "username": "mgranados", "puesto": "Oficial de Seguridad", "espacio": "Exteriores", "rol": "BASICO"}
]

# ==========================================
# GESTIÓN Y PERSISTENCIA DE API KEY GEMINI
# ==========================================
def obtener_gemini_api_key():
    try:
        conn = sqlite3.connect("database.db")
        c = conn.cursor()
        c.execute("CREATE TABLE IF NOT EXISTS config_sistema (clave TEXT PRIMARY KEY, valor TEXT)")
        c.execute("SELECT valor FROM config_sistema WHERE clave = 'gemini_api_key'")
        row = c.fetchone()
        conn.close()
        if row and row[0] and row[0].strip():
            return row[0].strip()
    except Exception:
        pass

    if hasattr(st, "secrets"):
        if "gemini_api_key" in st.secrets and str(st.secrets["gemini_api_key"]).strip():
            return str(st.secrets["gemini_api_key"]).strip()
        if "GEMINI_API_KEY" in st.secrets and str(st.secrets["GEMINI_API_KEY"]).strip():
            return str(st.secrets["GEMINI_API_KEY"]).strip()
        for sec in st.secrets:
            try:
                val = st.secrets[sec]
                if isinstance(val, dict) or hasattr(val, "get"):
                    for k in ["gemini_api_key", "GEMINI_API_KEY", "api_key", "key"]:
                        if k in val and str(val[k]).strip():
                            return str(val[k]).strip()
            except Exception:
                pass

    return os.environ.get("GEMINI_API_KEY", None)

def guardar_gemini_api_key(nueva_key):
    try:
        conn = sqlite3.connect("database.db")
        c = conn.cursor()
        c.execute("CREATE TABLE IF NOT EXISTS config_sistema (clave TEXT PRIMARY KEY, valor TEXT)")
        c.execute("INSERT OR REPLACE INTO config_sistema (clave, valor) VALUES ('gemini_api_key', ?)", (nueva_key.strip(),))
        conn.commit()
        conn.close()
        return True
    except Exception:
        return False

# ==========================================
# VERIFICACIÓN AUTOMÁTICA Y MIGRACIÓN BD
# ==========================================
def verificar_base_datos():
    if not os.path.exists("database.db"):
        instalar_sistema()
    else:
        conn = sqlite3.connect("database.db")
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='funcionarios'")
        tabla_existe = cursor.fetchone()
        conn.close()
        if not tabla_existe:
            instalar_sistema()

    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    
    cursor.execute("PRAGMA table_info(solicitudes)")
    cols_sol = [row[1] for row in cursor.fetchall()]
    nuevas_cols_sol = {
        'consecutivo': 'TEXT',
        'cantidad_despachada': 'INTEGER',
        'diferencia': 'INTEGER',
        'retirado_por': 'TEXT',
        'fecha_despacho_real': 'TEXT',
        'ruta_pdf_escaneado': 'TEXT'
    }
    for col, col_type in nuevas_cols_sol.items():
        if col not in cols_sol:
            try:
                cursor.execute(f"ALTER TABLE solicitudes ADD COLUMN {col} {col_type}")
            except Exception:
                pass

    cursor.execute("PRAGMA table_info(funcionarios)")
    cols_func = [row[1] for row in cursor.fetchall()]
    if 'debe_cambiar_password' not in cols_func:
        try:
            cursor.execute("ALTER TABLE funcionarios ADD COLUMN debe_cambiar_password INTEGER DEFAULT 0")
        except Exception:
            pass
    if 'email' not in cols_func:
        try:
            cursor.execute("ALTER TABLE funcionarios ADD COLUMN email TEXT")
        except Exception:
            pass

    hash_temp = hashlib.sha256(CLAVE_TEMPORAL_DEFAULT.encode()).hexdigest()
    for u in LISTA_USUARIOS_OFICIALES:
        cursor.execute("SELECT id FROM funcionarios WHERE username=?", (u["username"],))
        if not cursor.fetchone():
            cursor.execute("""
                INSERT INTO funcionarios (nombre, puesto, espacio_asignado, username, email, password, rol, estado, debe_cambiar_password)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'ACTIVO', 1)
            """, (u["nombre"], u["puesto"], u["espacio"], u["username"], u["email"], hash_temp, u["rol"]))

    conn.commit()
    conn.close()

verificar_base_datos()

# ===================================================================
# FUNCIÓN DE EXTRACCIÓN ROBUSTA DE PEDIMENTOS (DIGITAL O ESCANEO)
# ===================================================================
def extraer_datos_pedimento_pdf(pdf_bytes, filename):
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    full_text = ""
    for page in reader.pages:
        t = page.extract_text()
        if t:
            full_text += t + "\n"
            
    match_cons = re.search(r'Consecutivo:\s*([0-9]{3}-[0-9]{4})', full_text, re.IGNORECASE)
    if not match_cons:
        match_cons = re.search(r'([0-9]{3}-[0-9]{4})', filename)
    consecutivo = match_cons.group(1) if match_cons else "051-2026"
    
    articulos_encontrados = []
    
    pattern_bys = re.compile(r'(\d{5})\s+(\d{4,5})\s+(\d{5,6})\s+(\d{2})\s+(\d+)')
    matches_bys = pattern_bys.findall(full_text)
    
    if matches_bys:
        try:
            conn_cat = sqlite3.connect("database.db")
            c = conn_cat.cursor()
            for partida, subpartida, codigo, subcodigo, cantidad in matches_bys:
                c.execute("""
                    SELECT codigo_grupal, descripcion 
                    FROM articulos 
                    WHERE codigo_grupal LIKE ? AND codigo_grupal LIKE ?
                """, (f"{partida}%", f"%{codigo}{subcodigo}"))
                row = c.fetchone()
                if not row:
                    c.execute("""
                        SELECT codigo_grupal, descripcion 
                        FROM articulos 
                        WHERE codigo_grupal LIKE ? AND codigo_grupal LIKE ?
                    """, (f"{partida}%", f"%{codigo}%"))
                    row = c.fetchone()
                    
                if row:
                    articulos_encontrados.append({
                        "codigo": row[0],
                        "descripcion": row[1],
                        "cantidad_solicitada": int(cantidad),
                        "cantidad_despachada": int(cantidad)
                    })
            conn_cat.close()
        except Exception:
            pass

    if not articulos_encontrados:
        img_bytes = None
        for page in reader.pages:
            if page.images:
                img_bytes = page.images[0].data
                break
                
        api_key = obtener_gemini_api_key()
        if api_key and img_bytes:
            try:
                b64_img = base64.b64encode(img_bytes).decode("utf-8")
                prompt_ai = """
                Analiza este pedimento oficial de DINADECO. Extrae en JSON estricto:
                {
                  "consecutivo": "ej: 051-2026",
                  "articulos": [
                    {
                      "codigo": "Código numérico",
                      "descripcion": "Nombre del insumo",
                      "cantidad_solicitada": 4,
                      "cantidad_despachada": 4
                    }
                  ]
                }
                Devuelve ÚNICAMENTE el JSON.
                """
                payload_ai = {
                    "contents": [{
                        "parts": [
                            {"text": prompt_ai},
                            {"inlineData": {"mimeType": "image/jpeg", "data": b64_img}}
                        ]
                    }],
                    "generationConfig": {
                        "temperature": 0.1,
                        "responseMimeType": "application/json"
                    }
                }
                for mod in ["gemini-1.5-flash", "gemini-2.0-flash", "gemini-1.5-pro"]:
                    try:
                        url_ai = f"[https://generativelanguage.googleapis.com/v1beta/models/](https://generativelanguage.googleapis.com/v1beta/models/){mod}:generateContent?key={api_key}"
                        req_ai = urllib.request.Request(url_ai, data=json.dumps(payload_ai).encode('utf-8'), headers={'Content-Type': 'application/json'})
                        with urllib.request.urlopen(req_ai, timeout=20) as resp_ai:
                            data_ai = json.loads(resp_ai.read().decode('utf-8'))
                            raw_ai = data_ai['candidates'][0]['content']['parts'][0]['text']
                            raw_ai = re.sub(r'```json\s*', '', raw_ai).replace('```', '').strip()
                            parsed_ai = json.loads(raw_ai)
                            if parsed_ai.get("consecutivo"):
                                consecutivo = parsed_ai["consecutivo"]
                            items_ai = parsed_ai.get("articulos") or parsed_ai.get("items") or []
                            for a in items_ai:
                                articulos_encontrados.append({
                                    "codigo": a.get("codigo", ""),
                                    "descripcion": a.get("descripcion", ""),
                                    "cantidad_solicitada": int(a.get("cantidad_solicitada", 1)),
                                    "cantidad_despachada": int(a.get("cantidad_despachada", a.get("cantidad_solicitada", 1)))
                                })
                            break
                    except Exception:
                        continue
            except Exception:
                pass

    return consecutivo, articulos_encontrados

st.set_page_config(page_title="Bodega DINADECO", page_icon="📦", layout="wide", initial_sidebar_state="expanded")

if 'usuario_actual' not in st.session_state:
    st.session_state['usuario_actual'] = None
    st.session_state['rol'] = None
    st.session_state['func_id'] = None
if 'carrito' not in st.session_state:
    st.session_state['carrito'] = []
if 'ultimo_pedimento' not in st.session_state:
    st.session_state['ultimo_pedimento'] = None
if 'last_activity' not in st.session_state:
    st.session_state['last_activity'] = None
if 'debe_cambiar_password' not in st.session_state:
    st.session_state['debe_cambiar_password'] = False
if 'articulos_manuales_tab2' not in st.session_state:
    st.session_state['articulos_manuales_tab2'] = []

now = time.time()

def logout_inactividad():
    st.session_state['usuario_actual'] = None
    st.session_state['rol'] = None
    st.session_state['func_id'] = None
    st.session_state['carrito'] = []
    st.session_state['ultimo_pedimento'] = None
    st.session_state['last_activity'] = None
    st.session_state['debe_cambiar_password'] = False
    st.session_state['articulos_manuales_tab2'] = []
    st.query_params.clear()

def obtener_timestamp_query():
    if 't' in st.query_params:
        raw_t = st.query_params['t']
        if isinstance(raw_t, (list, tuple)):
            raw_t = raw_t[0]
        try:
            return float(raw_t)
        except (ValueError, TypeError):
            return None
    return None

last_t = obtener_timestamp_query()
last_active = st.session_state.get('last_activity') or last_t

if st.session_state['usuario_actual'] is None:
    if 'u' in st.query_params and last_t is not None:
        if (now - last_t) <= MAX_INACTIVIDAD:
            st.session_state['usuario_actual'] = st.query_params['u']
            st.session_state['rol'] = st.query_params.get('r', 'BASICO')
            st.session_state['func_id'] = int(st.query_params.get('fid', 0))
            st.session_state['last_activity'] = now
            st.query_params['t'] = str(int(now))
            
            conn = sqlite3.connect("database.db")
            cursor = conn.cursor()
            cursor.execute("SELECT COALESCE(debe_cambiar_password, 0) FROM funcionarios WHERE id=?", (st.session_state['func_id'],))
            row_cp = cursor.fetchone()
            conn.close()
            st.session_state['debe_cambiar_password'] = bool(row_cp[0]) if row_cp else False
        else:
            logout_inactividad()
            st.warning("⏱️ Su sesión ha expirado por inactividad. Inicie sesión de nuevo.")
    else:
        st.query_params.clear()
else:
    if last_active is not None and (now - last_active) > MAX_INACTIVIDAD:
        logout_inactividad()
        st.warning("⏱️ Su sesión se ha cerrado automáticamente por inactividad (5 minutos).")
        st.rerun()
    else:
        st.session_state['last_activity'] = now
        st.query_params['t'] = str(int(now))

def inyectar_temporizador_inactividad():
    js_codigo = """
    <script>
    (function() {
        const LIMITE_MS = 300000;
        let ultimaActividad = Date.now();
        let reloj;

        function cerrarSesionAuto() {
            try {
                window.top.location.href = window.top.location.origin + window.top.location.pathname;
            } catch(e) {
                window.location.href = window.location.origin + window.location.pathname;
            }
        }

        function registrarAccion() {
            const momento = Date.now();
            if (momento - ultimaActividad > LIMITE_MS) {
                cerrarSesionAuto();
                return;
            }
            ultimaActividad = momento;
            clearTimeout(reloj);
            reloj = setTimeout(cerrarSesionAuto, LIMITE_MS);
        }

        const eventos = ['mousemove', 'mousedown', 'keydown', 'touchstart', 'scroll', 'click'];
        eventos.forEach(ev => {
            try { window.top.addEventListener(ev, registrarAccion, true); } catch(e) {}
            try { window.addEventListener(ev, registrarAccion, true); } catch(e) {}
        });

        try { window.top.addEventListener('focus', registrarAccion); } catch(e) {}
        registrarAccion();
    })();
    </script>
    """
    components.html(js_codigo, height=0, width=0)

if st.session_state['rol'] != "ADMIN":
    st.markdown("""
        <style>
        [data-testid="stToolbarActions"],
        .stToolbarActionButton,
        [data-testid="stToolbar"] a,
        [data-testid="stToolbar"] button:not([kind="headerNoPadding"]),
        .stAppDeployButton,
        div[class*="viewerBadge"],
        #MainMenu,
        footer {
            display: none !important;
            visibility: hidden !important;
        }
        [data-testid="stSidebarCollapsedControl"] {
            display: block !important;
            visibility: visible !important;
        }
        [data-testid="stSidebarCollapsedControl"] button {
            display: inline-flex !important;
            visibility: visible !important;
        }
        </style>
    """, unsafe_allow_html=True)

def login(user, pwd):
    hash_pwd = hashlib.sha256(pwd.encode()).hexdigest()
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute("SELECT id, nombre, rol, estado, COALESCE(debe_cambiar_password, 0) FROM funcionarios WHERE username=? AND password=?", (user, hash_pwd))
    result = cursor.fetchone()
    conn.close()
    
    if result:
        if result[3] == "INACTIVO":
            st.error("Su cuenta está pendiente de aprobación por Jenier o Gabriela.")
        else:
            st.session_state['func_id'] = result[0]
            st.session_state['usuario_actual'] = result[1]
            st.session_state['rol'] = result[2]
            st.session_state['debe_cambiar_password'] = bool(result[4])
            st.session_state['last_activity'] = time.time()
            st.query_params['u'] = result[1]
            st.query_params['r'] = result[2]
            st.query_params['fid'] = str(result[0])
            st.query_params['t'] = str(int(time.time()))
            st.rerun()
    else:
        st.error("Usuario o contraseña incorrectos.")

def logout():
    logout_inactividad()
    st.rerun()

def sync_usuarios_excel():
    conn = sqlite3.connect("database.db")
    df_users = pd.read_sql_query("SELECT id, nombre, puesto, espacio_asignado, username, email, rol, estado FROM funcionarios", conn)
    conn.close()
    try:
        wb = openpyxl.load_workbook(EXCEL_MAESTRO)
        if 'USUARIOS' in wb.sheetnames:
            del wb['USUARIOS']
        ws = wb.create_sheet('USUARIOS')
        for r in dataframe_to_rows(df_users, index=False, header=True):
            ws.append(r)
        wb.save(EXCEL_MAESTRO)
    except Exception as e:
        st.error(f"Error sincronizando Excel Maestro: {e}")

def obtener_ultimo_consecutivo_db():
    anio_actual = datetime.now().year
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS control_consecutivos (
            anio INTEGER PRIMARY KEY,
            ultimo_numero INTEGER NOT NULL
        )
    """)
    cursor.execute("SELECT ultimo_numero FROM control_consecutivos WHERE anio = ?", (anio_actual,))
    row = cursor.fetchone()
    conn.close()
    return row[0] if row else 50

def fijar_ultimo_consecutivo_db(nuevo_num):
    anio_actual = datetime.now().year
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS control_consecutivos (
            anio INTEGER PRIMARY KEY,
            ultimo_numero INTEGER NOT NULL
        )
    """)
    cursor.execute("INSERT OR REPLACE INTO control_consecutivos (anio, ultimo_numero) VALUES (?, ?)", (anio_actual, nuevo_num))
    conn.commit()
    conn.close()

def limpiar_historial_y_archivos():
    conn = sqlite3.connect("database.db")
    cursor = conn.cursor()
    cursor.execute("DELETE FROM solicitudes")
    conn.commit()
    conn.close()

    for folder in ["pedimentos_generados", "pedimentos_escaneados"]:
        if os.path.exists(folder):
            archivos = glob.glob(os.path.join(folder, "*"))
            for f in archivos:
                try:
                    os.remove(f)
                except Exception:
                    pass

# -------------------------------------------------------------------
# PANTALLA DE LOGIN Y REGISTRO
# -------------------------------------------------------------------
if st.session_state['usuario_actual'] is None:
    tab_login, tab_registro = st.tabs(["🔑 Iniciar Sesión", "📝 Solicitar Acceso (Nuevo Usuario)"])
    
    with tab_login:
        st.subheader("Acceso al Sistema de Bodega")
        with st.form("form_login"):
            usuario = st.text_input("Usuario")
            clave = st.text_input("Contraseña", type="password")
            if st.form_submit_button("Ingresar"):
                login(usuario, clave)
                
    with tab_registro:
        st.subheader("Crear Nueva Cuenta")
        st.info("Su cuenta debe ser aprobada por jefatura (Jenier/Gabriela) antes de poder ingresar.")
        with st.form("form_registro"):
            n_nombre = st.text_input("Nombre Completo:")
            n_puesto = st.selectbox("Puesto / Rol:", ["Misceláneo/a", "Chofer", "Oficial de Seguridad", "Recepcionista", "Administrativo General"])
            n_espacio = st.selectbox("Espacio / Piso Asignado:", ["Piso 1 (30%)", "Piso 2 (40%)", "Piso 3 (30%)", "Oficina General", "Exteriores", "Flota Vehicular"])
            n_user = st.text_input("Nombre de Usuario (ej. cruiz):")
            n_pass = st.text_input("Contraseña:", type="password")
            
            if st.form_submit_button("Solicitar Acceso"):
                if n_nombre and n_user and n_pass:
                    conn = sqlite3.connect("database.db")
                    cursor = conn.cursor()
                    cursor.execute("SELECT id FROM funcionarios WHERE username=?", (n_user,))
                    if cursor.fetchone():
                        st.error("Ese nombre de usuario ya existe.")
                    else:
                        hash_p = hashlib.sha256(n_pass.encode()).hexdigest()
                        cursor.execute("INSERT INTO funcionarios (nombre, puesto, espacio_asignado, username, password, rol, estado, debe_cambiar_password) VALUES (?, ?, ?, ?, ?, 'BASICO', 'INACTIVO', 0)", 
                                       (n_nombre, n_puesto, n_espacio, n_user, hash_p))
                        conn.commit()
                        conn.close()
                        sync_usuarios_excel()
                        st.success("Cuenta registrada. Jenier o Gabriela deben aprobar el ingreso.")
                else:
                    st.error("Por favor complete todos los campos.")
    st.stop()

# ===================================================================
# PANTALLA OBLIGATORIA: CAMBIO DE CONTRASEÑA EN EL PRIMER INICIO
# ===================================================================
if st.session_state.get('debe_cambiar_password', False):
    st.warning("🔒 **Primer Inicio de Sesión:** Por motivos de seguridad institucional, debe definir una contraseña personal antes de continuar.")
    with st.form("form_cambio_clave_obligatorio"):
        st.write(f"Usuario: **{st.session_state['usuario_actual']}**")
        nueva_clave = st.text_input("Nueva Contraseña:", type="password")
        conf_clave = st.text_input("Confirmar Nueva Contraseña:", type="password")
        
        col_c1, col_c2 = st.columns([1, 4])
        if col_c1.form_submit_button("💾 Guardar y Entrar", type="primary"):
            if not nueva_clave or len(nueva_clave) < 6:
                st.error("La nueva contraseña debe tener al menos 6 caracteres.")
            elif nueva_clave != conf_clave:
                st.error("Las contraseñas no coinciden.")
            elif nueva_clave == CLAVE_TEMPORAL_DEFAULT:
                st.error("No puede utilizar la misma contraseña temporal provisional.")
            else:
                hash_nueva = hashlib.sha256(nueva_clave.encode()).hexdigest()
                conn = sqlite3.connect("database.db")
                cursor = conn.cursor()
                cursor.execute("UPDATE funcionarios SET password=?, debe_cambiar_password=0 WHERE id=?", (hash_nueva, st.session_state['func_id']))
                conn.commit()
                conn.close()
                sync_usuarios_excel()
                st.session_state['debe_cambiar_password'] = False
                st.success("✅ ¡Contraseña actualizada exitosamente! Ingresando al sistema...")
                time.sleep(1)
                st.rerun()
        if col_c2.form_submit_button("Cerrar Sesión"):
            logout()
    st.stop()

inyectar_temporizador_inactividad()

# -------------------------------------------------------------------
# NAVEGACIÓN Y BOTÓN DE CIERRE DE SESIÓN EN SIDEBAR
# -------------------------------------------------------------------
st.sidebar.markdown("### 👤 Sesión Activa")
st.sidebar.write(f"Usuario: **{st.session_state['usuario_actual']}**")
st.sidebar.write(f"Rol: **{st.session_state['rol']}**")

if st.sidebar.button("🔴 Cerrar Sesión", type="primary", use_container_width=True, key="btn_logout_sidebar"):
    logout()

st.sidebar.markdown("---")

if st.session_state['rol'] == "ADMIN":
    tabs = st.tabs([
        "📝 1. Pedimentos", 
        "✍️ 2. Firma/Aprobación PDF", 
        "📦 3. Despacho/Recepción Bodega", 
        "📊 4. Informes", 
        "⏱️ 5. Ciclos de Entrega y Vida Útil",
        "✅ 6. Aprobar Usuarios", 
        "⚙️ 7. Control Master"
    ])
    tab_pedimentos, tab_firma, tab_recepcion, tab_informes, tab_frecuencias, tab_usuarios, tab_master = tabs
else:
    tabs = st.tabs(["📝 1. Mis Pedimentos (Hasta 10 ítems)"])
    tab_pedimentos = tabs[0]

# -------------------------------------------------------------------
# TAB 1: FORMULARIO DE PEDIMENTOS (CON CANDADO DE VIDA ÚTIL)
# -------------------------------------------------------------------
with tab_pedimentos:
    col_u1, col_u2 = st.columns([4, 1])
    with col_u1:
        st.caption(f"👤 Conectado: **{st.session_state['usuario_actual']}** ({st.session_state['rol']})")
    with col_u2:
        if st.button("🔴 Salir", key="btn_logout_top", use_container_width=True):
            logout()

    st.header("📦 Solicitud de Insumos")
    
    if st.session_state['ultimo_pedimento']:
        ped = st.session_state['ultimo_pedimento']
        st.success(f"🎉 ¡Pedimento procesado con éxito! Consecutivo asignado: **#{ped['consecutivo']}**")
        
        if ped.get("exito_mail"):
            st.info("📩 El documento ha sido enviado por correo para la respectiva Firma Digital.")
        else:
            st.error(f"⚠️ **Atención:** Ocurrió un inconveniente al enviar el correo automático: **{ped.get('msg_mail')}**")
        st.markdown("---")

    conn = sqlite3.connect("database.db")
    df_articulos = pd.read_sql_query("SELECT codigo_grupal, descripcion, categoria, frecuencia_entrega, dias_vida_util FROM articulos", conn)
    
    cursor = conn.cursor()
    cursor.execute("SELECT espacio_asignado FROM funcionarios WHERE id=?", (st.session_state['func_id'],))
    res_piso = cursor.fetchone()
    piso_asignado = res_piso[0] if res_piso else "Oficina General"

    colA, colB = st.columns([1, 2])
    
    with colA:
        st.subheader("Añadir Insumos")
        categorias = ["Todas"] + list(df_articulos["categoria"].dropna().unique())
        cat_sel = st.selectbox("Categoría:", categorias)
        
        df_filtro = df_articulos if cat_sel == "Todas" else df_articulos[df_articulos["categoria"] == cat_sel]
        articulo_sel = st.selectbox("Artículo:", df_filtro["descripcion"].tolist())
        item_info = df_articulos[df_articulos["descripcion"] == articulo_sel].iloc[0]
        
        cantidad = st.number_input("Cantidad:", min_value=1, value=1, step=1)
        
        cursor.execute("""
            SELECT COALESCE(fecha_despacho_real, fecha_aprobacion)
            FROM solicitudes
            WHERE funcionario_id = ? AND articulo_id = ? AND estado LIKE 'ENTREGADO%'
            ORDER BY id DESC LIMIT 1
        """, (st.session_state['func_id'], item_info['codigo_grupal']))
        last_req = cursor.fetchone()
        
        bloqueado = False
        msg_bloqueo = ""
        dias_faltantes = 0
        
        if last_req and last_req[0] and pd.notna(item_info['dias_vida_util']):
            try:
                fecha_str = str(last_req[0]).split()[0]
                if "/" in fecha_str:
                    last_date = datetime.strptime(fecha_str, "%d/%m/%Y").date()
                else:
                    last_date = datetime.strptime(fecha_str, "%Y-%m-%d").date()
                    
                dias_transcurridos = (datetime.now().date() - last_date).days
                vida_util = int(item_info['dias_vida_util'])
                
                if dias_transcurridos < vida_util:
                    bloqueado = True
                    dias_faltantes = vida_util - dias_transcurridos
                    msg_bloqueo = f"⛔ **Frecuencia de Entrega No Cumplida:**\nEste artículo se le despachó físicamente el **{last_date.strftime('%d/%m/%Y')}**.\nDebe esperar **{dias_faltantes} días** más para volver a solicitarlo (Vida útil reglamentaria: {vida_util} días)."
            except Exception:
                pass
                
        if bloqueado:
            st.error(msg_bloqueo)
            st.info("💡 Si ocurrió una emergencia (ej. derrame, daño, rotura), justifíquelo para habilitar la solicitud por excepción:")
            just_excepcion = st.text_input("Justificación de Excepción:", key="just_exc")
            if just_excepcion.strip():
                if st.button("⚠️ Solicitar por Excepción"):
                    if len(st.session_state['carrito']) >= 10:
                        st.error("Límite máximo de 10 ítems por pedimento alcanzado.")
                    elif any(item['codigo'] == item_info['codigo_grupal'] for item in st.session_state['carrito']):
                        st.warning("Este artículo ya está en su lista actual.")
                    else:
                        st.session_state['carrito'].append({
                            "codigo": item_info['codigo_grupal'],
                            "descripcion": f"{articulo_sel} [EXCEPCIÓN: {just_excepcion}]",
                            "cantidad": cantidad,
                            "frecuencia": item_info['frecuencia_entrega'],
                            "dias_vida": item_info['dias_vida_util']
                        })
                        st.success("Artículo agregado al pedimento como excepción.")
        else:
            if st.button("➕ Agregar al Pedimento"):
                if len(st.session_state['carrito']) >= 10:
                    st.error("Límite máximo de 10 ítems por pedimento alcanzado.")
                elif any(item['codigo'] == item_info['codigo_grupal'] for item in st.session_state['carrito']):
                    st.warning("Este artículo ya está en su lista actual.")
                else:
                    st.session_state['carrito'].append({
                        "codigo": item_info['codigo_grupal'],
                        "descripcion": articulo_sel,
                        "cantidad": cantidad,
                        "frecuencia": item_info['frecuencia_entrega'],
                        "dias_vida": item_info['dias_vida_util']
                    })
                    st.success("Artículo agregado al pedimento.")
        
    conn.close()

    with colB:
        st.subheader(f"🛒 Mi Pedimento Actual ({len(st.session_state['carrito'])}/10)")
        if len(st.session_state['carrito']) > 0:
            df_carrito = pd.DataFrame(st.session_state['carrito'])
            st.dataframe(df_carrito[['descripcion', 'cantidad', 'frecuencia']], use_container_width=True)
            
            if st.button("🗑️ Vaciar Lista"):
                st.session_state['carrito'] = []
                st.rerun()
            
            justificacion_global = st.text_area("Justificación general del pedido (Opcional):")
            if st.button("🚀 Enviar Pedimento Completo a Jefatura", type="primary"):
                conn = sqlite3.connect("database.db")
                cursor = conn.cursor()
                
                ids_guardados = []
                for item in st.session_state['carrito']:
                    cursor.execute("""
                        INSERT INTO solicitudes (funcionario_id, articulo_id, cantidad, es_valido, justificacion, estado)
                        VALUES (?, ?, ?, ?, ?, 'PENDIENTE')
                    """, (st.session_state['func_id'], item['codigo'], item['cantidad'], 1, justificacion_global))
                    ids_guardados.append(cursor.lastrowid)
                conn.commit()
                
                hash_temp = hashlib.sha256(str(datetime.now()).encode()).hexdigest()
                consecutivo, ruta_excel = generar_formulario_oficial_bys(ids_guardados, "Departamento de Servicios Generales B", "Jenier Camacho Picado", hash_temp)
                
                placeholders = ",".join("?" * len(ids_guardados))
                cursor.execute(f"UPDATE solicitudes SET consecutivo = ? WHERE id IN ({placeholders})", [consecutivo] + ids_guardados)
                conn.commit()
                
                ruta_pdf = ruta_excel.replace(".xlsx", ".pdf")
                exportar_excel_a_pdf(ruta_excel, ruta_pdf)
                conn.close()

                exito_mail, msg_mail = enviar_correo_con_pdf_adjunto(
                    consecutivo=consecutivo, 
                    lista_items=st.session_state['carrito'], 
                    solicitante=st.session_state['usuario_actual'], 
                    es_valido=True, 
                    motivo="Pedido reglamentario", 
                    ruta_pdf=ruta_pdf, 
                    justificacion=justificacion_global
                )

                st.session_state['ultimo_pedimento'] = {
                    "consecutivo": consecutivo,
                    "exito_mail": exito_mail,
                    "msg_mail": msg_mail
                }
                st.session_state['carrito'] = []
                st.rerun()
        else:
            st.info("No ha agregado artículos a la lista.")

# -------------------------------------------------------------------
# SECCIONES EXCLUSIVAS PARA ADMINISTRADORES
# -------------------------------------------------------------------
if st.session_state['rol'] == "ADMIN":
    
    # ===============================================================
    # TAB 2: FIRMA Y APROBACIÓN
    # ===============================================================
    with tab_firma:
        st.header("✍️ Validación de Pedimentos Firmados en Adobe Acrobat")
        
        st.subheader("1. Pedimentos Pendientes de Firma")
        st.write("Descargue el documento PDF que desea firmar en Adobe Acrobat:")
        
        if os.path.exists("pedimentos_generados"):
            archivos_pdf = [f for f in os.listdir("pedimentos_generados") if f.endswith(".pdf")]
            if archivos_pdf:
                for arch in archivos_pdf:
                    ruta_c = os.path.join("pedimentos_generados", arch)
                    with open(ruta_c, "rb") as f_arch:
                        st.download_button(
                            label=f
